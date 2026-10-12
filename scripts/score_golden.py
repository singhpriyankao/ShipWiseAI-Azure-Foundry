"""
Compares the evaluators with the reviewer's verdicts on the golden cases.

Usage:
    python score_golden.py [--judge "path/to/foundry results.jsonl"] [--export-dataset out.jsonl]

Reads evals/golden_hhh_cases.json, the saved runs in evals/golden_runs/ and the reviewer's verdicts in
evals/review_sheet.xlsx (columns J to L of the Review sheet).

Two evaluators are compared with the reviewer:
  code   the deterministic checks named in the golden file (golden_checks.py); a case passes when all its checks pass
  judge  a Foundry prompt-based judge's label, from a results .jsonl exported from a Foundry evaluation run
A judge row that carries an error (for example 'Invalid or missing numeric result', or sample.error) is counted as
'evaluator failed' and left out of the agreement figures: Foundry marks such a row completed and failed, which would
otherwise look like a real fail.

Reviewer verdicts of 'unsure' are left out of the agreement figures. With a handful of rows per agent these are counts, not
rates to quote as precision; the output prints the rows so each disagreement can be read.

--export-dataset writes the 19 rows (query, response, ground_truth, agents_run, id) in the format the Foundry evaluation
takes; the ground truth is the reviewer's edited expected behaviour if there is one, otherwise the drafted one.
"""
import collections
import json
import os
import sys

from openpyxl import load_workbook

import golden_checks

HERE = os.path.dirname(os.path.abspath(__file__))
CASES_FILE = os.path.join(HERE, "..", "evals", "golden_hhh_cases.json")
RUNS_DIR = os.path.join(HERE, "..", "evals", "golden_runs")
SHEET = os.path.join(HERE, "..", "evals", "review_sheet.xlsx")
AGENT = {"regulations": "Regulations", "competitive_analysis": "Competitive", "model_stack_analysis": "Model Stack"}


def load_review():
    ws = load_workbook(SHEET)["Review"]
    out = {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        out[r[0]] = {"verdict": (r[9] or "").strip().lower() or None, "note": r[10], "edited": r[11], "expected": r[7]}
    return out


def load_judge(path):
    """{case id: 'pass' | 'fail' | 'error'} from a Foundry results .jsonl; judge rows are found by evaluator name."""
    out = {}
    for line in open(path, encoding="utf-8"):
        d = json.loads(line)
        cid = (d.get("datasource_item") or {}).get("id")
        for r in d.get("results") or []:
            if "judge" not in (r.get("name") or "").lower():
                continue
            err = (r.get("sample") or {}).get("error")
            reason = r.get("reason") or ""
            if err or reason.startswith("Please provide a valid evaluator") or r.get("status") not in ("completed", None) or not r.get("label"):
                out[cid] = "error"
            else:
                out[cid] = "pass" if r.get("passed") else "fail"
    return out


def main():
    args = sys.argv[1:]
    cases = json.load(open(CASES_FILE, encoding="utf-8"))["cases"]
    review = load_review()
    judge = load_judge(args[args.index("--judge") + 1]) if "--judge" in args else {}
    rows = []
    for case in cases:
        path = os.path.join(RUNS_DIR, f"{case['id']}.json")
        if not os.path.exists(path):
            continue
        record = json.load(open(path, encoding="utf-8"))["record"]
        checks = golden_checks.run_checks(case, record)
        rows.append({"case": case, "record": record, "checks": checks, "code": "pass" if all(p for _, p, _ in checks) else "fail",
                     "routing": golden_checks.routing_ok(case, record), "human": review.get(case["id"], {}).get("verdict"),
                     "judge": judge.get(case["id"])})

    if "--export-dataset" in args:
        out = args[args.index("--export-dataset") + 1]
        with open(out, "w", encoding="utf-8") as f:
            for r in rows:
                rv = review.get(r["case"]["id"], {})
                truth = (rv.get("edited") or rv.get("expected") or "").strip()
                f.write(json.dumps({"id": r["case"]["id"], "query": r["case"]["message"], "response": r["record"].get("reply") or "",
                                    "ground_truth": truth, "agents_run": ",".join(r["record"].get("agents_run") or []) or "none"}, ensure_ascii=False) + "\n")
        print(f"{len(rows)} rows written to {out}")
        return

    print(f"{'case':18} {'reviewer':9} {'code':6} {'judge':7} routing")
    for r in rows:
        print(f"{r['case']['id']:18} {str(r['human']):9} {r['code']:6} {str(r['judge'] or '-'):7} {'as expected' if r['routing'] else 'different from expected'}")
    print()
    for who in ("code", "judge"):
        if who == "judge" and not judge:
            continue
        print(f"== {who} against the reviewer (rows with an 'unsure' verdict or an evaluator error are left out)")
        by_agent = collections.defaultdict(lambda: collections.Counter())
        for r in rows:
            v = r[who]
            if r["human"] not in ("pass", "fail") or v not in ("pass", "fail"):
                by_agent[AGENT[r["case"]["component"]]]["left out"] += 1
                continue
            by_agent[AGENT[r["case"]["component"]]]["agree" if v == r["human"] else ("missed a fail" if r["human"] == "fail" else "false alarm")] += 1
        for agent, c in sorted(by_agent.items()):
            n = c["agree"] + c["missed a fail"] + c["false alarm"]
            print(f"  {agent:12} {c['agree']} of {n} agree; {c['missed a fail']} missed a fail (said pass, reviewer said fail); {c['false alarm']} false alarms (said fail, reviewer said pass); {c['left out']} left out")
        print()
    print("== disagreements, to read one by one")
    for r in rows:
        for who in ("code", "judge"):
            v = r[who]
            if v in ("pass", "fail") and r["human"] in ("pass", "fail") and v != r["human"]:
                why = "; ".join(f"{c['type']}: {n}" for c, p, n in r["checks"] if not p) if who == "code" else ""
                print(f"  {r['case']['id']}: {who} said {v}, reviewer said {r['human']}" + (f" ({why})" if why else "") + f"\n     reviewer note: {review[r['case']['id']]['note']}")
    odd = [r["case"]["id"] for r in rows if not r["routing"]]
    if odd:
        print(f"\n== routing differed from the case's expected agents: {', '.join(odd)}")


if __name__ == "__main__":
    main()
